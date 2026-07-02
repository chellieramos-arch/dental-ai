"""
DentAI — Image Index Builder
─────────────────────────────
One-time (and safely re-runnable/resumable) batch job that builds a dedicated,
independently-searchable image index — separate from the text index — so
image relevance is judged by what an image actually shows, not by which page
happened to answer the question in words.

For every PDF/PPTX in ./documents, this:
  1. Extracts candidate images using the same size/aspect/margin/decorative
     filters already used at query time in app.py.
  2. Sends each candidate to Claude (vision) once, asking it to (a) write a
     caption describing exactly what's depicted and (b) classify whether it's
     a genuine clinical diagram/photo/chart worth showing students, versus
     decorative art, a logo, a cover page, or a screenshot of typed text.
  3. Uploads the image bytes to Supabase Storage (bucket 'course-images') and
     embeds the caption with the same embedding model used for text, upserting
     into Pinecone tagged content_type="image" — so it can be searched at
     query time directly against the student's question, independent of the
     text-chunk retrieval that answers the question in words.

Only images classified as genuine diagrams get indexed. Everything else is
counted as "skipped" and never stored — keeping the index clean by construction
rather than filtering junk out reactively on every query.

SCALE WARNING: a corpus this size is expected to have on the order of tens of
thousands of candidate images. This will take hours and make that many Claude
vision calls (small thumbnails, cheap per-call, but it adds up) plus one
OpenAI embedding call per kept image. Progress is checkpointed per-file to
image_index_progress.json, so it is safe to Ctrl+C and re-run later — already
-completed files are skipped, not re-billed.

Run:
    MODE=cloud python3 build_image_index.py
"""

import os
import io
import json
import time
import hashlib
import sys

from dotenv import load_dotenv
load_dotenv()

from config import IS_CLOUD, PINECONE_API_KEY, PINECONE_INDEX, SUPABASE_URL, SUPABASE_SERVICE_KEY, SUPABASE_KEY

PDF_FOLDER          = "./documents"
PROGRESS_FILE       = "image_index_progress.json"
IMAGE_BUCKET        = "course-images"
EMBED_MODEL         = "text-embedding-ada-002"
CLAUDE_VISION_MODEL = "claude-haiku-4-5-20251001"   # cheap — this is a per-image classification call, not a full clinical answer
MIN_PX              = 120
UPSERT_BATCH        = 50

if not IS_CLOUD:
    sys.exit("This builds the CLOUD (Pinecone) image index. Run with: MODE=cloud python3 build_image_index.py")


# ── Progress checkpoint — safe to interrupt and resume ────────────────────────
def load_progress():
    if os.path.exists(PROGRESS_FILE):
        with open(PROGRESS_FILE) as f:
            return json.load(f)
    return {"done_files": [], "stats": {"images_indexed": 0, "images_skipped": 0, "files_done": 0}}


def save_progress(progress):
    with open(PROGRESS_FILE, "w") as f:
        json.dump(progress, f, indent=2)


# ── Junk filtering — mirrors extract_page_images() in app.py ─────────────────
def _looks_decorative(image_bytes: bytes) -> bool:
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(image_bytes)).convert("L")
        im.thumbnail((64, 64))
        pixels = list(im.getdata())
        if not pixels:
            return False
        mean = sum(pixels) / len(pixels)
        variance = sum((p - mean) ** 2 for p in pixels) / len(pixels)
        return (variance ** 0.5) < 10
    except Exception:
        return False


def extract_candidates_from_pdf(path, min_px=MIN_PX):
    import fitz
    candidates = []
    try:
        doc = fitz.open(path)
    except Exception:
        return candidates
    for page_idx, page in enumerate(doc):
        page_area = max(page.rect.width * page.rect.height, 1)
        for img_info in page.get_images(full=True):
            try:
                base = doc.extract_image(img_info[0])
            except Exception:
                continue
            w, h = base["width"], base["height"]
            if w < min_px or h < min_px:
                continue
            if max(w, h) / max(min(w, h), 1) > 8:
                continue
            try:
                bbox = page.get_image_bbox(img_info)
                margin = page.rect.height * 0.08
                if bbox.y1 < margin or bbox.y0 > page.rect.height - margin:
                    continue
                if (bbox.width * bbox.height) / page_area < 0.02:
                    continue
            except Exception:
                pass
            if _looks_decorative(base["image"]):
                continue
            candidates.append({"bytes": base["image"], "ext": base["ext"],
                                "page": page_idx + 1, "w": w, "h": h})
    doc.close()
    return candidates


def extract_candidates_from_pptx(path, min_px=MIN_PX):
    from pptx import Presentation
    candidates = []
    try:
        prs = Presentation(path)
    except Exception:
        return candidates
    for slide_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.shape_type != 13:   # MSO_SHAPE_TYPE.PICTURE
                continue
            try:
                w, h = shape.image.size
                if w < min_px or h < min_px:
                    continue
                if max(w, h) / max(min(w, h), 1) > 8:
                    continue
                if _looks_decorative(shape.image.blob):
                    continue
                candidates.append({"bytes": shape.image.blob, "ext": shape.image.ext,
                                    "page": slide_idx + 1, "w": w, "h": h})
            except Exception:
                pass
    return candidates


# ── Caption + classify via Claude vision ──────────────────────────────────────
def caption_image(anthropic_client, image_bytes, file_name, page_num, max_retries=4):
    import base64 as b64
    from PIL import Image as PILImage

    try:
        im = PILImage.open(io.BytesIO(image_bytes))
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        im.thumbnail((500, 500))
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=75)
        data = b64.standard_b64encode(buf.getvalue()).decode("utf-8")
    except Exception as e:
        return None, f"thumbnail failed: {e}"

    for attempt in range(max_retries):
        try:
            resp = anthropic_client.messages.create(
                model=CLAUDE_VISION_MODEL,
                max_tokens=250,
                tools=[{
                    "name": "describe_image",
                    "description": "Describe a candidate image from dental course materials and classify it.",
                    "input_schema": {
                        "type": "object",
                        "properties": {
                            "is_diagram": {
                                "type": "boolean",
                                "description": (
                                    "True if this is a genuine clinical diagram, anatomical "
                                    "illustration, clinical/radiograph photo, chart, or step-by-step "
                                    "figure. False if it's decorative art, a logo, a cover/title "
                                    "graphic, a screenshot of typed text/table, or otherwise not "
                                    "useful to show a student."
                                ),
                            },
                            "caption": {
                                "type": "string",
                                "description": (
                                    "One or two sentences describing exactly what is depicted — "
                                    "specific enough that someone could match it to a clinical "
                                    "question (procedure, tooth/anatomy, view type, what's labeled). "
                                    "If is_diagram is False, briefly say why (e.g. 'logo', 'text screenshot')."
                                ),
                            },
                        },
                        "required": ["is_diagram", "caption"],
                    },
                }],
                tool_choice={"type": "tool", "name": "describe_image"},
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": (
                            f"This image is from '{file_name}', page/slide {page_num}, in a dental "
                            f"school's course materials. Describe and classify it."
                        )},
                        {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": data}},
                    ],
                }],
            )
            for block in resp.content:
                if block.type == "tool_use" and block.name == "describe_image":
                    return block.input, None
            return None, "no tool_use block in response"
        except Exception as e:
            is_overloaded = getattr(e, "status_code", None) == 529 or "overloaded" in str(e).lower()
            if not is_overloaded or attempt == max_retries - 1:
                return None, f"{type(e).__name__}: {e}"
            time.sleep(2 ** attempt)
    return None, "exhausted retries"


def main():
    import anthropic
    import openai
    from pinecone import Pinecone
    from supabase import create_client

    anthropic_client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"), max_retries=3)
    oai   = openai.OpenAI()
    pc    = Pinecone(api_key=PINECONE_API_KEY)
    index = pc.Index(PINECONE_INDEX)

    sb_key = SUPABASE_SERVICE_KEY or SUPABASE_KEY
    if not sb_key:
        sys.exit("No SUPABASE_SERVICE_KEY or SUPABASE_KEY set in .env")
    sb = create_client(SUPABASE_URL, sb_key)
    try:
        sb.storage.create_bucket(IMAGE_BUCKET, options={"public": False})
        print(f"[image-index] Created Supabase bucket '{IMAGE_BUCKET}'.")
    except Exception:
        pass  # already exists — fine. If uploads fail below, create it manually in the Supabase dashboard.

    progress   = load_progress()
    done_files = set(progress["done_files"])

    files = sorted(f for f in os.listdir(PDF_FOLDER) if f.lower().endswith((".pdf", ".pptx", ".ppt")))
    files = [f for f in files if f not in done_files]
    print(f"[image-index] {len(files)} file(s) left to process ({len(done_files)} already done this run/previous runs).")
    print(f"[image-index] So far: {progress['stats']['images_indexed']} indexed, {progress['stats']['images_skipped']} skipped.\n")

    for fi, fname in enumerate(files):
        path = os.path.join(PDF_FOLDER, fname)
        print(f"[{fi + 1}/{len(files)}] {fname}", flush=True)

        if fname.lower().endswith(".pdf"):
            candidates = extract_candidates_from_pdf(path)
        else:
            candidates = extract_candidates_from_pptx(path)

        vectors_batch = []
        for c in candidates:
            result, err = caption_image(anthropic_client, c["bytes"], fname, c["page"])
            if err or not result:
                progress["stats"]["images_skipped"] += 1
                continue
            if not result.get("is_diagram"):
                progress["stats"]["images_skipped"] += 1
                continue

            caption = (result.get("caption") or "")[:800]
            storage_key = (
                f"{hashlib.md5(fname.encode()).hexdigest()[:10]}_p{c['page']}_"
                f"{hashlib.md5(c['bytes']).hexdigest()[:8]}.{c['ext']}"
            )
            try:
                try:
                    sb.storage.from_(IMAGE_BUCKET).remove([storage_key])
                except Exception:
                    pass
                sb.storage.from_(IMAGE_BUCKET).upload(
                    storage_key, c["bytes"], {"content-type": f"image/{c['ext']}"}
                )
            except Exception as e:
                print(f"    ⚠️  upload failed for {storage_key}: {e}")
                progress["stats"]["images_skipped"] += 1
                continue

            try:
                emb = oai.embeddings.create(input=caption, model=EMBED_MODEL).data[0].embedding
            except Exception as e:
                print(f"    ⚠️  embedding failed: {e}")
                progress["stats"]["images_skipped"] += 1
                continue

            vectors_batch.append({
                "id": f"img::{fname}::{c['page']}::{storage_key}",
                "values": emb,
                "metadata": {
                    "content_type": "image",
                    "file_name":    fname,
                    "page_label":   str(c["page"]),
                    "caption":      caption,
                    "storage_key":  storage_key,
                    "width":        c["w"],
                    "height":       c["h"],
                },
            })
            progress["stats"]["images_indexed"] += 1

            if len(vectors_batch) >= UPSERT_BATCH:
                index.upsert(vectors=vectors_batch)
                vectors_batch = []

        if vectors_batch:
            index.upsert(vectors=vectors_batch)

        progress["done_files"].append(fname)
        progress["stats"]["files_done"] += 1
        save_progress(progress)   # checkpoint after every file — safe to Ctrl+C between files
        print(f"    ✓ {len(candidates)} candidate(s) checked — "
              f"{progress['stats']['images_indexed']} indexed / "
              f"{progress['stats']['images_skipped']} skipped so far\n")

    print(f"\n[image-index] Done! {progress['stats']['images_indexed']} images indexed, "
          f"{progress['stats']['images_skipped']} skipped, across {progress['stats']['files_done']} files.")


if __name__ == "__main__":
    main()
